"""Rule engine.

Two consumers, one definition:

* :func:`evaluate_mask` -- pure pandas.  Used for the HTML report, the flag
  sheet and for rules Excel cannot express (stuck value, sigma outlier, jump).
* :func:`excel_rules` -- turns a :class:`Rule` into real openpyxl
  *conditional formatting* objects, so the colours stay editable in Excel and
  keep working when the user sorts or filters.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Any, Iterable, Optional

import numpy as np
import pandas as pd

from .models import ColumnSpec, Rule, RuleKind

# Rules that must be pre-computed in Python (no Excel equivalent).
STATIC_KINDS = {RuleKind.OUTLIER_SIGMA, RuleKind.STUCK, RuleKind.JUMP}
# Rules that paint a gradient rather than a flag.
GRADIENT_KINDS = {RuleKind.COLOR_SCALE, RuleKind.DATA_BAR}


# --------------------------------------------------------------------------- #
# python evaluation
# --------------------------------------------------------------------------- #
def evaluate_mask(series: pd.Series, rule: Rule) -> pd.Series:
    """Boolean mask of the cells this rule flags.  Gradients return all-False."""
    kind = rule.kind
    n = len(series)
    false = pd.Series(np.zeros(n, dtype=bool), index=series.index)
    if kind in GRADIENT_KINDS or n == 0:
        return false

    numeric = pd.api.types.is_numeric_dtype(series)

    if kind is RuleKind.BLANK:
        if numeric:
            return series.isna()
        return series.isna() | (series.astype("string").fillna("").str.strip() == "")

    if kind is RuleKind.TEXT_CONTAINS:
        if not rule.text:
            return false
        return (series.astype("string").fillna("")
                .str.contains(rule.text, case=False, regex=False))

    if kind is RuleKind.DUPLICATE:
        return series.duplicated(keep=False) & series.notna()

    if not numeric:
        return false

    s = series.astype("float64")
    valid = s.notna()

    if kind is RuleKind.ZERO:
        tol = abs(float(rule.tolerance or 0.0))
        return valid & (s.abs() <= tol)

    if kind is RuleKind.NEGATIVE:
        return valid & (s < 0)

    if kind is RuleKind.THRESHOLD:
        v = float(rule.value)
        op = rule.operator
        table = {
            "<": s < v, "<=": s <= v, ">": s > v, ">=": s >= v,
            "==": s == v, "!=": s != v,
        }
        return valid & table.get(op, s < v)

    if kind is RuleKind.BETWEEN:
        lo, hi = sorted((float(rule.value), float(rule.value2)))
        return valid & s.between(lo, hi)

    if kind is RuleKind.TOP_BOTTOM:
        vals = s.dropna()
        if vals.empty:
            return false
        k = max(1, int(round(len(vals) * rule.rank / 100.0)) if rule.percent else int(rule.rank))
        k = min(k, len(vals))
        cutoff = (vals.nsmallest(k).max() if rule.bottom else vals.nlargest(k).min())
        return valid & ((s <= cutoff) if rule.bottom else (s >= cutoff))

    if kind is RuleKind.OUTLIER_SIGMA:
        k = float(rule.value2 or 3.0)
        mu, sd = s.mean(), s.std(ddof=0)
        if not np.isfinite(sd) or sd == 0:
            return false
        return valid & ((s - mu).abs() > k * sd)

    if kind is RuleKind.STUCK:
        window = max(2, int(rule.value2 or 10))
        # groups of consecutive identical values
        grp = (s != s.shift()).cumsum()
        sizes = s.groupby(grp).transform("size")
        return valid & (sizes >= window)

    if kind is RuleKind.JUMP:
        limit = abs(float(rule.value2 or rule.value or 0.0))
        if limit == 0:
            return false
        return valid & (s.diff().abs() > limit)

    return false


def hex_color(color: str) -> str:
    """``"FF184F95"`` / ``"184F95"`` / ``"#184F95"`` -> ``"#184F95"``."""
    c = str(color or "").lstrip("#")
    if len(c) == 8:            # openpyxl ARGB
        c = c[2:]
    return "#" + c.upper() if c else ""


def ink_for(color: str) -> str:
    """Readable text colour on ``color``: near-black or white, by WCAG contrast.

    Used by the Excel band fills and by the HTML table so both label a given
    fill the same way.  A brightness average would put white on mid-tone reds
    and greens where black reads much better.
    """
    c = hex_color(color).lstrip("#")
    if len(c) != 6:
        return "#111111"
    def lin(v: int) -> float:
        x = v / 255.0
        return x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4
    r, g, b = (int(c[i:i + 2], 16) for i in (0, 2, 4))
    lum = 0.2126 * lin(r) + 0.7152 * lin(g) + 0.0722 * lin(b)
    vs_dark = (lum + 0.05) / (0.00561 + 0.05)      # against #111111
    vs_light = 1.05 / (lum + 0.05)
    return "#111111" if vs_dark >= vs_light else "#FFFFFF"


def excel_argb(color: str) -> str:
    """``"#184F95"`` -> ``"FF184F95"`` for openpyxl."""
    c = str(color or "").lstrip("#")
    if len(c) == 6:
        return "FF" + c.upper()
    if len(c) == 8:
        return c.upper()
    return "FFFFFFFF"


def gradient_style(rules: Iterable[Rule], series: Optional[pd.Series] = None
                   ) -> dict[str, Any]:
    """The colour scheme every output should use: Excel, HTML and PNG alike.

    Reads the first enabled colour-gradient rule.  Without one, falls back to
    a plain 2-colour blue ramp so nothing is left undefined.
    """
    rules = list(rules)
    rule = next((r for r in rules
                 if r.kind is RuleKind.COLOR_SCALE and r.enabled), None)
    zero = next((r for r in rules if r.kind is RuleKind.ZERO and r.enabled), None)
    bar = next((r for r in rules if r.kind is RuleKind.DATA_BAR and r.enabled), None)
    common = {
        "zero_color": (hex_color(zero.fill) if zero is not None else ""),
        "zero_label": (zero.display_name() if zero is not None else ""),
        "bars": bar is not None,
        "bar_color": (hex_color(bar.fill) if bar is not None else "#4a90d9"),
    }
    if rule is None:
        return {
            "mode": "2color",
            "colors": ["#FCFCFB", "#184F95"],
            "bounds": "auto",
            "stops": [],
            "reverse": False,
            "from_rule": False,
            "label": "",
            **common,
        }
    return {
        "mode": rule.scale_mode,
        "colors": [hex_color(c) for c in band_colors(rule)],
        "bounds": rule.scale_bounds,
        "stops": [float(v) for v in (scale_stops(rule, series) or [])],
        "reverse": bool(rule.scale_reverse),
        "from_rule": True,
        "label": rule.display_name(),
        **common,
    }


def active_rules(rules: Iterable[Rule], column: str, is_numeric: bool) -> list[Rule]:
    return [r for r in rules if r.applies_to(column, is_numeric)]


@dataclass
class Flag:
    row: int
    time: Any
    entity: str
    column: str
    label: str
    rule: str
    value: Any


def collect_flags(
    df: pd.DataFrame,
    specs: dict[str, ColumnSpec],
    rules: list[Rule],
    time_column: Optional[str],
    entity: str,
    limit: int = 20000,
) -> list[Flag]:
    """Every rule hit as a flat list - becomes the "flags" sheet."""
    out: list[Flag] = []
    for col in df.columns:
        if col == time_column:
            continue
        spec = specs.get(col)
        if spec is not None and not spec.include:
            continue
        is_num = pd.api.types.is_numeric_dtype(df[col])
        for rule in active_rules(rules, col, is_num):
            if rule.kind in GRADIENT_KINDS:
                continue
            mask = evaluate_mask(df[col], rule)
            if not mask.any():
                continue
            idx = list(np.flatnonzero(mask.to_numpy()))
            for i in idx:
                if len(out) >= limit:
                    return out
                out.append(
                    Flag(
                        row=int(i) + 1,
                        time=(df[time_column].iloc[i] if time_column in df.columns else ""),
                        entity=entity,
                        column=(spec.name if spec and spec.name else col),
                        label=rule.display_name(),
                        rule=rule.kind.value,
                        value=df[col].iloc[i],
                    )
                )
    return out


# --------------------------------------------------------------------------- #
# openpyxl conditional formatting
# --------------------------------------------------------------------------- #
# Five-band default palette: low -> high (red, orange, yellow, light green, green).
BAND_FILLS = ["FFF8696B", "FFFCAA78", "FFFFEB84", "FFB1D580", "FF63BE7B"]


def scale_stops(rule: Rule, series: Optional[pd.Series]) -> list[float]:
    """Numeric boundaries for a gradient / band rule, low to high."""
    n_stops = {"2color": 2, "3color": 3, "5color": 5}.get(rule.scale_mode, 3)
    if rule.scale_bounds == "numbers" and rule.scale_values:
        vals = sorted(float(v) for v in rule.scale_values)
        return vals[:n_stops] if len(vals) >= n_stops else vals

    data = None
    if series is not None and pd.api.types.is_numeric_dtype(series):
        data = series.astype("float64").dropna()
    if data is None or data.empty:
        return list(np.linspace(0.0, 1.0, n_stops))

    if rule.scale_bounds == "percent":
        qs = np.linspace(0.05, 0.95, n_stops)
        return [float(data.quantile(q)) for q in qs]
    lo, hi = float(data.min()), float(data.max())
    if lo == hi:
        hi = lo + 1.0
    return list(np.linspace(lo, hi, n_stops))


def band_colors(rule: Rule) -> list[str]:
    """Colours for the chosen gradient mode, low to high."""
    n = {"2color": 2, "3color": 3, "5color": 5}.get(rule.scale_mode, 3)
    colors = [c for c in rule.scale_colors if c]
    if len(colors) < n:
        if n == 2:
            colors = [colors[0] if colors else BAND_FILLS[0], BAND_FILLS[-1]]
        elif n == 3:
            base = colors + BAND_FILLS[:3]
            colors = [base[0], base[1], base[2]]
        else:
            colors = list(BAND_FILLS)
    colors = colors[:n]
    if rule.scale_reverse:
        colors = colors[::-1]
    return colors


def _color_scale_rules(rule: Rule, series: Optional[pd.Series]) -> list[Any]:
    """2/3-colour gradient, or 5 discrete bands (Excel gradients stop at 3)."""
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
    from openpyxl.styles import Font, PatternFill

    colors = band_colors(rule)
    stops = scale_stops(rule, series)
    explicit = rule.scale_bounds == "numbers" or (
        rule.scale_bounds == "percent" and series is not None)

    if rule.scale_mode == "5color":
        # 5 solid bands: <s1, s1..s2, s2..s3, s3..s4, >=s4
        if len(stops) < 5:
            stops = list(np.linspace(stops[0], stops[-1], 5)) if len(stops) >= 2 else \
                [0, 0.25, 0.5, 0.75, 1.0]
        out = []
        edges = stops[1:]                       # 4 boundaries -> 5 bands
        for i, color in enumerate(colors):
            fill = PatternFill("solid", start_color=color, end_color=color)
            font = Font(color=rule.font)
            if i == 0:
                out.append(CellIsRule(operator="lessThan", formula=[repr(float(edges[0]))],
                                      fill=fill, font=font))
            elif i == len(colors) - 1:
                out.append(CellIsRule(operator="greaterThanOrEqual",
                                      formula=[repr(float(edges[-1]))], fill=fill, font=font))
            else:
                lo, hi = float(edges[i - 1]), float(edges[i])
                out.append(CellIsRule(operator="between",
                                      formula=[repr(lo), repr(hi)], fill=fill, font=font))
        return out

    if len(colors) >= 3:
        if explicit and len(stops) >= 3:
            return [ColorScaleRule(
                start_type="num", start_value=stops[0], start_color=colors[0],
                mid_type="num", mid_value=stops[len(stops) // 2], mid_color=colors[1],
                end_type="num", end_value=stops[-1], end_color=colors[2])]
        return [ColorScaleRule(
            start_type="min", start_color=colors[0],
            mid_type="percentile", mid_value=50, mid_color=colors[1],
            end_type="max", end_color=colors[2])]

    if explicit and len(stops) >= 2:
        return [ColorScaleRule(start_type="num", start_value=stops[0], start_color=colors[0],
                               end_type="num", end_value=stops[-1], end_color=colors[-1])]
    return [ColorScaleRule(start_type="min", start_color=colors[0],
                           end_type="max", end_color=colors[-1])]


def excel_rules(rule: Rule, series: Optional[pd.Series] = None) -> list[Any]:
    """openpyxl rule objects for one :class:`Rule` (empty for static kinds)."""
    from openpyxl.formatting.rule import (
        CellIsRule,
        ColorScaleRule,
        DataBarRule,
        FormulaRule,
        Rule as XLRule,
    )
    from openpyxl.styles import Font, PatternFill
    from openpyxl.styles.differential import DifferentialStyle

    if rule.kind in STATIC_KINDS:
        return []

    fill = PatternFill(start_color=rule.fill, end_color=rule.fill, fill_type="solid")
    font = Font(color=rule.font, bold=rule.bold)
    kind = rule.kind

    if kind is RuleKind.ZERO:
        tol = abs(float(rule.tolerance or 0.0))
        if tol:
            return [CellIsRule(operator="between", formula=[f"-{tol}", f"{tol}"],
                               fill=fill, font=font, stopIfTrue=False)]
        return [CellIsRule(operator="equal", formula=["0"], fill=fill, font=font,
                           stopIfTrue=False)]

    if kind is RuleKind.NEGATIVE:
        return [CellIsRule(operator="lessThan", formula=["0"], fill=fill, font=font)]

    if kind is RuleKind.BLANK:
        return [XLRule(type="containsBlanks",
                       dxf=DifferentialStyle(fill=fill, font=font), stopIfTrue=False)]

    if kind is RuleKind.THRESHOLD:
        ops = {"<": "lessThan", "<=": "lessThanOrEqual", ">": "greaterThan",
               ">=": "greaterThanOrEqual", "==": "equal", "!=": "notEqual"}
        return [CellIsRule(operator=ops.get(rule.operator, "lessThan"),
                           formula=[repr(float(rule.value))], fill=fill, font=font)]

    if kind is RuleKind.BETWEEN:
        lo, hi = sorted((float(rule.value), float(rule.value2)))
        return [CellIsRule(operator="between", formula=[repr(lo), repr(hi)],
                           fill=fill, font=font)]

    if kind is RuleKind.DUPLICATE:
        return [XLRule(type="duplicateValues",
                       dxf=DifferentialStyle(fill=fill, font=font))]

    if kind is RuleKind.TEXT_CONTAINS:
        return [XLRule(type="containsText", operator="containsText", text=rule.text,
                       dxf=DifferentialStyle(fill=fill, font=font))]

    if kind is RuleKind.TOP_BOTTOM:
        return [XLRule(type="top10", rank=int(rule.rank), percent=bool(rule.percent),
                       bottom=bool(rule.bottom),
                       dxf=DifferentialStyle(fill=fill, font=font))]

    if kind is RuleKind.COLOR_SCALE:
        return _color_scale_rules(rule, series)

    if kind is RuleKind.DATA_BAR:
        return [DataBarRule(start_type="min", end_type="max", color=rule.bar_color,
                            showValue=True, minLength=None, maxLength=None)]

    return []
