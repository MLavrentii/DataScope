"""Excel output: translated, filterable, conditionally formatted workbooks.

Two products:

* :func:`write_cleaned_workbook` - one workbook per source file
  (``<name>_cleaned.xlsx``).  Original files are never touched.
* :func:`write_analysis_workbook` - one workbook comparing all entities:
  time x PCS matrices, deviation from the fleet, daily summary, zero report.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Optional, Sequence

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.marker import Marker
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .analysis import AnalysisResult, MatrixResult
from .models import ColumnSpec, OutputOptions, Rule, RuleKind
from .rules import (
    GRADIENT_KINDS,
    STATIC_KINDS,
    active_rules,
    evaluate_mask,
    excel_argb,
    excel_rules,
    gradient_style,
    ink_for,
)

# ---- house style ---------------------------------------------------------- #
HDR_FILL = PatternFill("solid", start_color="FF1F3A5F", end_color="FF1F3A5F")
HDR_FONT = Font(color="FFFFFFFF", bold=True, size=10)
SUB_FILL = PatternFill("solid", start_color="FFEDEDEA", end_color="FFEDEDEA")
SUB_FONT = Font(color="FF6B6B66", size=8, italic=True)
TITLE_FONT = Font(bold=True, size=12, color="FF1F3A5F")
NOTE_FONT = Font(size=9, color="FF52514E")
THIN = Side(style="thin", color="FFD9D9D6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TIME_FMT = "yyyy-mm-dd hh:mm:ss"

SERIES_COLORS = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4",
                 "008300", "4A3AA7", "E34948"]
MAX_STATIC_FILLS = 60000
MAX_CHART_POINTS = 1500


@dataclass
class SheetWriteResult:
    sheet: str
    rows: int
    cf_rules: int
    static_cells: int


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def _num_format(spec: Optional[ColumnSpec], series: Optional[pd.Series]) -> str:
    if spec is not None and spec.decimals is not None:
        return "0" if spec.decimals == 0 else "0." + "0" * int(spec.decimals)
    if series is None or not pd.api.types.is_numeric_dtype(series):
        return "General"
    s = series.dropna()
    if s.empty:
        return "General"
    if (s == s.round(0)).all():
        return "#,##0"
    mx = float(s.abs().max())
    if mx >= 1000:
        return "#,##0.0"
    if mx >= 10:
        return "0.00"
    return "0.000"


def _autosize(ws: Worksheet, headers: Sequence[str], default: int = 14,
              first_width: int = 20) -> None:
    for i, h in enumerate(headers, start=1):
        text = str(h).split("\n")[0]
        width = max(len(text) * 1.35 + 2, 8)
        # CJK characters are roughly double width
        cjk = sum(1 for ch in text if ord(ch) > 0x2E80)
        width += cjk * 0.9
        ws.column_dimensions[get_column_letter(i)].width = min(
            max(default, width), 34 if i > 1 else max(first_width, 20)
        )


def _write_table(
    ws: Worksheet,
    df: pd.DataFrame,
    headers: Sequence[str],
    start_row: int = 1,
    sub_headers: Optional[Sequence[str]] = None,
    time_column: Optional[str] = None,
    number_formats: Optional[dict[str, str]] = None,
    freeze: bool = True,
    autofilter: bool = True,
) -> int:
    """Write a DataFrame with a styled header.  Returns the first data row."""
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    row = start_row + 1
    if sub_headers:
        for j, h in enumerate(sub_headers, start=1):
            c = ws.cell(row=row, column=j, value=h)
            c.fill = SUB_FILL
            c.font = SUB_FONT
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
            c.border = BOX
        row += 1
    data_start = row

    values = df.to_numpy(dtype=object, na_value=None) if len(df) else np.empty((0, len(headers)),
                                                                              dtype=object)
    for r in values:
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in r])

    # formats
    nf = number_formats or {}
    for j, col in enumerate(df.columns, start=1):
        letter = get_column_letter(j)
        fmt = TIME_FMT if col == time_column else nf.get(str(col))
        if not fmt:
            continue
        for cell in ws[f"{letter}{data_start}:{letter}{max(data_start, ws.max_row)}"]:
            cell[0].number_format = fmt

    ws.row_dimensions[start_row].height = 32
    if freeze:
        ws.freeze_panes = ws.cell(row=data_start, column=2)
    if autofilter and len(df):
        ws.auto_filter.ref = (f"A{start_row}:"
                             f"{get_column_letter(len(headers))}{ws.max_row}")
    _autosize(ws, headers)
    return data_start


def _apply_rules_to_sheet(
    ws: Worksheet,
    df: pd.DataFrame,
    specs: dict[str, ColumnSpec],
    rules: list[Rule],
    data_start: int,
    time_column: Optional[str] = None,
    header_row: int = 1,
    aliases: Optional[dict[str, list[str]]] = None,
    max_cf_columns: int = 0,
) -> tuple[int, int]:
    """Conditional formatting + static fills.  Returns (cf_count, static_cells).

    ``aliases`` lets a column be matched by more than one name - a per-unit
    sheet holds ``Data01`` but the rule may have been written against
    ``PCS01_Data01``.
    """
    cf_count = 0
    static_cells = 0
    cf_columns = 0
    last_row = data_start + len(df) - 1
    if last_row < data_start:
        return 0, 0

    for j, col in enumerate(df.columns, start=1):
        if col == time_column:
            continue
        spec = specs.get(str(col))
        if spec is not None and not spec.include:
            continue
        if max_cf_columns and cf_columns >= max_cf_columns:
            break
        series = df[col]
        is_num = pd.api.types.is_numeric_dtype(series)
        letter = get_column_letter(j)
        rng = f"{letter}{data_start}:{letter}{last_row}"
        names = [str(col)] + list((aliases or {}).get(str(col), []))
        matched = [r for r in rules if any(r.applies_to(n, is_num) for n in names)]
        if matched:
            cf_columns += 1
        for rule in matched:
            if rule.kind in STATIC_KINDS:
                if static_cells >= MAX_STATIC_FILLS:
                    continue
                mask = evaluate_mask(series, rule)
                if not mask.any():
                    continue
                fill = PatternFill("solid", start_color=rule.fill, end_color=rule.fill)
                font = Font(color=rule.font, bold=rule.bold)
                for i in np.flatnonzero(mask.to_numpy()):
                    if static_cells >= MAX_STATIC_FILLS:
                        break
                    cell = ws.cell(row=data_start + int(i), column=j)
                    cell.fill = fill
                    cell.font = font
                    static_cells += 1
                continue
            if rule.kind in GRADIENT_KINDS and not is_num:
                continue
            for xr in excel_rules(rule, series):
                ws.conditional_formatting.add(rng, xr)
                cf_count += 1
    return cf_count, static_cells


def _line_chart(
    ws_data: Worksheet,
    ws_target: Worksheet,
    anchor: str,
    title: str,
    y_title: str,
    min_col: int,
    max_col: int,
    min_row: int,
    max_row: int,
    cat_col: int = 1,
    width: float = 30,
    height: float = 12,
    markers: bool = False,
) -> None:
    """A restrained line chart: 2px lines, no 3-D, recessive gridlines."""
    ch = LineChart()
    ch.title = title
    ch.style = 2
    ch.height = height
    ch.width = width
    ch.y_axis.title = y_title
    ch.x_axis.title = None
    data = Reference(ws_data, min_col=min_col, max_col=max_col,
                     min_row=min_row - 1, max_row=max_row)
    cats = Reference(ws_data, min_col=cat_col, max_col=cat_col,
                     min_row=min_row, max_row=max_row)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    for i, s in enumerate(ch.series):
        color = SERIES_COLORS[i % len(SERIES_COLORS)]
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = 20000        # ~2 pt
        s.smooth = False
        if markers:
            s.marker = Marker(symbol="circle", size=5)
        else:
            s.marker = Marker(symbol="none")
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ws_target.add_chart(ch, anchor)


def _downsample(df: pd.DataFrame, limit: int = MAX_CHART_POINTS) -> pd.DataFrame:
    if len(df) <= limit:
        return df
    step = int(np.ceil(len(df) / limit))
    return df.iloc[::step]


# --------------------------------------------------------------------------- #
# per-file workbook
# --------------------------------------------------------------------------- #
def write_cleaned_workbook(
    out_path: Path,
    df: pd.DataFrame,
    specs: dict[str, ColumnSpec],
    rules: list[Rule],
    opts: OutputOptions,
    time_column: Optional[str],
    info: dict[str, Any],
    flags: Optional[pd.DataFrame] = None,
    chart_metrics: Optional[list[str]] = None,
    header_template: str = "{name}\n[{unit}]",
    keep_raw_row: bool = True,
    png_paths: Optional[list[Path]] = None,
) -> Path:
    """One translated + coloured workbook for one source file."""
    keep = [c for c in df.columns
            if (specs.get(str(c)) is None or specs[str(c)].include)]
    df = df[keep]

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    headers, subs, nfmt = [], [], {}
    for c in df.columns:
        spec = specs.get(str(c))
        if c == time_column:
            headers.append(spec.name if spec and spec.name else str(c))
        elif spec is not None:
            headers.append(spec.header(header_template))
        else:
            headers.append(str(c))
        subs.append(str(c))
        nfmt[str(c)] = _num_format(spec, df[c] if c != time_column else None)

    data_start = _write_table(
        ws, df, headers, start_row=1,
        sub_headers=subs if keep_raw_row else None,
        time_column=time_column, number_formats=nfmt,
        freeze=opts.freeze_panes, autofilter=opts.autofilter,
    )
    cf_count, static_cells = _apply_rules_to_sheet(
        ws, df, specs, rules, data_start, time_column
    )

    # ---- Info sheet ----------------------------------------------------- #
    wsi = wb.create_sheet("Info")
    wsi["A1"] = "DataScope - file report"
    wsi["A1"].font = TITLE_FONT
    r = 3
    for k, v in info.items():
        wsi.cell(row=r, column=1, value=str(k)).font = Font(bold=True, size=10)
        wsi.cell(row=r, column=2, value=("" if v is None else str(v))).font = NOTE_FONT
        r += 1
    r += 1
    wsi.cell(row=r, column=1, value="Rules applied").font = TITLE_FONT
    r += 1
    for rule in rules:
        if not rule.enabled:
            continue
        scope = rule.scope.value
        if rule.scope.value == "selected":
            scope = f"selected ({len(rule.columns)} col.)"
        elif rule.scope.value == "pattern":
            scope = f"pattern /{rule.pattern}/"
        wsi.cell(row=r, column=1, value=rule.display_name()).font = Font(size=10)
        wsi.cell(row=r, column=2, value=scope).font = NOTE_FONT
        cell = wsi.cell(row=r, column=3, value="   ")
        if rule.kind not in GRADIENT_KINDS:
            cell.fill = PatternFill("solid", start_color=rule.fill, end_color=rule.fill)
        r += 1
    r += 1
    unmatched = info.get("Untranslated columns") or ""
    if unmatched:
        wsi.cell(row=r, column=1, value="Columns not found in the dictionary").font = TITLE_FONT
        r += 1
        for name in str(unmatched).split(" | "):
            wsi.cell(row=r, column=1, value=name).font = NOTE_FONT
            r += 1
    wsi.column_dimensions["A"].width = 38
    wsi.column_dimensions["B"].width = 60
    wsi.column_dimensions["C"].width = 8

    # ---- Flags sheet ---------------------------------------------------- #
    if flags is not None and not flags.empty:
        wsf = wb.create_sheet("Flags")
        _write_table(wsf, flags, list(flags.columns), start_row=1,
                     time_column="Time",
                     number_formats={"Time": TIME_FMT},
                     freeze=True, autofilter=True)

    # ---- Charts --------------------------------------------------------- #
    if opts.excel_charts and chart_metrics and time_column is not None:
        metrics = [m for m in chart_metrics if m in df.columns][:8]
        if metrics:
            wsc = wb.create_sheet("Charts")
            small = _downsample(df[[time_column] + metrics])
            hdrs = [str(specs[m].name if specs.get(m) and specs[m].name else m) for m in metrics]
            wsd = wb.create_sheet("ChartData")
            wsd.sheet_state = "hidden"
            wsd.append(["Time"] + hdrs)
            for _, row in small.iterrows():
                wsd.append([row[time_column]] + [None if pd.isna(row[m]) else float(row[m])
                                                 for m in metrics])
            for i, m in enumerate(metrics):
                spec = specs.get(m)
                _line_chart(
                    wsd, wsc, anchor=f"A{2 + i * 24}",
                    title=(spec.name if spec and spec.name else m),
                    y_title=(spec.unit if spec else ""),
                    min_col=2 + i, max_col=2 + i,
                    min_row=2, max_row=wsd.max_row,
                )
            if png_paths:
                col = "T"
                for i, p in enumerate(png_paths):
                    try:
                        img = XLImage(str(p))
                        wsc.add_image(img, f"{col}{2 + i * 24}")
                    except Exception:
                        pass

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


# --------------------------------------------------------------------------- #
# per-unit workbook (wide source files)
# --------------------------------------------------------------------------- #
def write_entity_workbook(
    out_path: Path,
    table: Any,                       # loader.LoadedTable
    layout: Any,                      # wide.WideLayout
    specs: dict[str, ColumnSpec],     # raw column -> spec
    mspecs: dict[str, ColumnSpec],    # metric key  -> spec (name without prefix)
    rules: list[Rule],
    opts: OutputOptions,
    info: dict[str, Any],
    alarms: Optional[pd.DataFrame] = None,
    header_template: str = "{name}\n[{unit}]",
    keep_raw_row: bool = True,
    chart_metrics: Optional[list[str]] = None,
) -> Path:
    """One sheet per unit: 16 readable columns instead of 689 cryptic ones."""
    from .wide import split_tables

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Info"
    ws0["A1"] = "DataScope - file report"
    ws0["A1"].font = TITLE_FONT
    r = 3
    for k, v in info.items():
        ws0.cell(row=r, column=1, value=str(k)).font = Font(bold=True, size=10)
        ws0.cell(row=r, column=2, value=("" if v is None else str(v))).font = NOTE_FONT
        r += 1
    r += 1
    ws0.cell(row=r, column=1, value="Rules applied").font = TITLE_FONT
    r += 1
    for rule in rules:
        if not rule.enabled:
            continue
        scope = rule.scope.value
        if scope == "selected":
            scope = f"selected ({len(rule.columns)} col.)"
        elif scope == "pattern":
            scope = f"pattern /{rule.pattern}/"
        ws0.cell(row=r, column=1, value=rule.display_name()).font = Font(size=10)
        ws0.cell(row=r, column=2, value=scope).font = NOTE_FONT
        cell = ws0.cell(row=r, column=3, value="   ")
        if rule.kind not in GRADIENT_KINDS:
            cell.fill = PatternFill("solid", start_color=rule.fill, end_color=rule.fill)
        r += 1
    ws0.column_dimensions["A"].width = 34
    ws0.column_dimensions["B"].width = 62
    ws0.column_dimensions["C"].width = 8

    time_col = table.time_column
    subs = split_tables(table, layout)
    for sub in subs:
        ws = wb.create_sheet(str(sub.entity)[:31])
        cols = ([time_col] if time_col in sub.df.columns else []) + [
            c for c in sub.df.columns if c != time_col
            and (mspecs.get(str(c)) is None or mspecs[str(c)].include)
        ]
        df = sub.df.loc[:, cols]
        headers, raws, nfmt, aliases = [], [], {}, {}
        for c in df.columns:
            spec = mspecs.get(str(c))
            if c == time_col:
                headers.append("時刻 / Time")
            elif spec is not None:
                headers.append(spec.header(header_template))
            else:
                headers.append(str(c))
            raw = layout.cells.get((sub.entity, str(c)), str(c))
            raws.append(raw)
            aliases[str(c)] = [raw]
            nfmt[str(c)] = _num_format(spec, df[c] if c != time_col else None)
        data_start = _write_table(
            ws, df, headers, start_row=1,
            sub_headers=raws if keep_raw_row else None,
            time_column=time_col, number_formats=nfmt,
            freeze=opts.freeze_panes, autofilter=opts.autofilter,
        )
        _apply_rules_to_sheet(ws, df, mspecs, rules, data_start, time_col,
                              aliases=aliases, max_cf_columns=opts.max_cf_columns)

    # plant-level columns (irradiance, weather, receiving panel, curtailment)
    plant_cols = [c for c in layout.plant_columns if c in table.df.columns]
    if plant_cols:
        ws = wb.create_sheet("PLANT")
        cols = ([time_col] if time_col in table.df.columns else []) + plant_cols
        df = table.df.loc[:, cols]
        headers, raws, nfmt = [], [], {}
        for c in df.columns:
            spec = specs.get(str(c))
            headers.append("時刻 / Time" if c == time_col else
                           (spec.header(header_template) if spec else str(c)))
            raws.append(str(c))
            nfmt[str(c)] = _num_format(spec, df[c] if c != time_col else None)
        data_start = _write_table(ws, df, headers, start_row=1,
                                  sub_headers=raws if keep_raw_row else None,
                                  time_column=time_col, number_formats=nfmt,
                                  freeze=opts.freeze_panes, autofilter=opts.autofilter)
        _apply_rules_to_sheet(ws, df, specs, rules, data_start, time_col,
                              max_cf_columns=opts.max_cf_columns)

    if alarms is not None and not alarms.empty:
        wsa = wb.create_sheet("STATUS_ALARMS")
        start = _write_table(wsa, alarms, list(alarms.columns), start_row=1,
                             freeze=True, autofilter=True)
        if "Share %" in alarms.columns:
            j = list(alarms.columns).index("Share %") + 1
            letter = get_column_letter(j)
            gpal = [excel_argb(c) for c in
                    (gradient_style(rules).get("colors") or ["FFFCFCFB", "FF184F95"])]
            wsa.conditional_formatting.add(
                f"{letter}{start}:{letter}{start + len(alarms) - 1}",
                ColorScaleRule(start_type="num", start_value=0, start_color=gpal[0],
                               end_type="num", end_value=100, end_color=gpal[-1]))

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


# --------------------------------------------------------------------------- #
# analysis workbook
# --------------------------------------------------------------------------- #
def _gradient_cf(style: dict[str, Any], diverging: bool,
                 values: Optional[pd.Series] = None) -> list[Any]:
    """Conditional formatting for a matrix range, built from the app's own
    colour-scale setting (see :func:`core.rules.gradient_style`).

    The analysis sheets used to be hard-coded blue, which contradicted whatever
    the user had configured; now Excel, the HTML report and the PNGs all read
    the same style dict.
    """
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Font, PatternFill

    colors = [excel_argb(c) for c in (style.get("colors") or ["FFFCFCFB", "FF184F95"])]
    if len(colors) < 2:
        colors = [colors[0] if colors else "FFFCFCFB", "FF184F95"]
    lo, mid, hi = colors[0], colors[len(colors) // 2], colors[-1]

    if diverging:
        # deviation: low colour for below reference, high colour for above
        return [ColorScaleRule(
            start_type="percentile", start_value=2, start_color=lo,
            mid_type="num", mid_value=0, mid_color="FFFFFFFF",
            end_type="percentile", end_value=98, end_color=hi)]

    stops = [float(v) for v in (style.get("stops") or [])]
    fixed = style.get("bounds") == "numbers" and len(stops) >= 2

    if len(colors) >= 5 and fixed:
        # five solid bands - Excel's own gradients stop at three colours
        edges = list(np.linspace(stops[0], stops[-1], 5))[1:] if len(stops) < 5 else stops[1:5]
        out: list[Any] = []
        for i, color in enumerate(colors[:5]):
            fill = PatternFill("solid", start_color=color, end_color=color)
            font = Font(color=excel_argb(ink_for(color)))
            if i == 0:
                out.append(CellIsRule(operator="lessThan", formula=[repr(float(edges[0]))],
                                      fill=fill, font=font))
            elif i == 4:
                out.append(CellIsRule(operator="greaterThanOrEqual",
                                      formula=[repr(float(edges[3]))], fill=fill, font=font))
            else:
                out.append(CellIsRule(operator="between",
                                      formula=[repr(float(edges[i - 1])), repr(float(edges[i]))],
                                      fill=fill, font=font))
        return out

    if fixed:
        n = len(stops)
        return [ColorScaleRule(
            start_type="num", start_value=stops[0], start_color=lo,
            mid_type="num", mid_value=stops[n // 2], mid_color=mid,
            end_type="num", end_value=stops[-1], end_color=hi)]

    if len(colors) == 2:
        return [ColorScaleRule(start_type="min", start_color=lo,
                               end_type="max", end_color=hi)]
    return [ColorScaleRule(start_type="min", start_color=lo,
                           mid_type="percentile", mid_value=50, mid_color=mid,
                           end_type="max", end_color=hi)]


PICK_SHEET = "_pick"
PICK_FILL = PatternFill("solid", start_color="FFEDEDEA", end_color="FFEDEDEA")
OUT_OF_RANGE_FONT = Font(color="FFBFBEB8", size=10)


def _pick_lists(wb: Workbook, dates: Sequence[str]) -> tuple[str, str]:
    """Hidden sheet holding the dropdown choices; returns the two ranges.

    Inline validation lists are capped at 255 characters, which 18 dates plus
    24 hours would blow past, so the choices live on a hidden sheet.
    """
    ws = wb[PICK_SHEET] if PICK_SHEET in wb.sheetnames else wb.create_sheet(PICK_SHEET)
    ws.sheet_state = "hidden"
    if ws.max_row < 2:                      # fill it once per workbook
        ws["A1"] = "dates"
        ws["B1"] = "hours"
        for i, d in enumerate(dates, start=2):
            c = ws.cell(row=i, column=1, value=pd.Timestamp(d).to_pydatetime())
            c.number_format = "yyyy-mm-dd"
        for h in range(24):
            ws.cell(row=2 + h, column=2, value=h)
    n = max(2, len(dates) + 1)
    return (f"{PICK_SHEET}!$A$2:$A${n}", f"{PICK_SHEET}!$B$2:$B$25")


def _period_picker(ws: Worksheet, wb: Workbook, matrix: pd.DataFrame,
                   data_start: int, last_row: int, flag_col: int) -> None:
    """A click-to-choose period on a matrix sheet.

    A plain .xlsx has no calendar control (that needs VBA and a 32-bit ActiveX
    control), so this is the closest macro-free equivalent: dropdowns listing
    the dates and hours that are actually in the data, an "in range" flag
    column, and greyed-out rows outside the window.
    """
    from openpyxl.worksheet.datavalidation import DataValidation

    idx = pd.DatetimeIndex(matrix.index)
    if not len(idx):
        return
    dates = sorted({str(d) for d in idx.normalize().strftime("%Y-%m-%d")})
    date_range, hour_range = _pick_lists(wb, dates)

    ws["A2"] = "期間 / Period"
    ws["A2"].font = Font(bold=True, size=10, color="FF1F3A5F")
    # short labels: row 2 shares its columns with the data table (~14 wide),
    # so a long label would spill over the cell next to it
    labels = {"B2": "日付/Date", "D2": "開始/From", "F2": "終了/To"}
    for cell, text in labels.items():
        ws[cell] = text
        ws[cell].font = NOTE_FONT
    ws["C2"] = pd.Timestamp(dates[0]).to_pydatetime()
    ws["C2"].number_format = "yyyy-mm-dd"
    ws["E2"] = int(idx.min().hour)
    ws["G2"] = int(idx.max().hour)
    for cell in ("C2", "E2", "G2"):
        ws[cell].fill = PICK_FILL
        ws[cell].border = BOX
        ws[cell].font = Font(bold=True, size=10)

    dv_date = DataValidation(type="list", formula1=f"={date_range}", allow_blank=True)
    dv_hour = DataValidation(type="list", formula1=f"={hour_range}", allow_blank=True)
    ws.add_data_validation(dv_date)
    ws.add_data_validation(dv_hour)
    dv_date.add(ws["C2"])
    dv_hour.add(ws["E2"])
    dv_hour.add(ws["G2"])

    # the flag column: 1 while the timestamp is inside the chosen window
    letter = get_column_letter(flag_col)
    ws.cell(row=3, column=flag_col, value="範囲内 / In range").fill = HDR_FILL
    ws.cell(row=3, column=flag_col).font = HDR_FONT
    ws.cell(row=3, column=flag_col).alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True)
    for r in range(data_start, last_row + 1):
        ws.cell(row=r, column=flag_col, value=(
            f'=IF(AND($A{r}>=$C$2+$E$2/24,$A{r}<$C$2+($G$2+1)/24),1,0)'))
    ws.column_dimensions[letter].width = 15
    # "2026-08-06 00:00:00" needs more than the default first-column width,
    # otherwise Excel shows ###
    ws.column_dimensions["A"].width = 22

    # grey everything outside the window - instant feedback, no filtering needed
    note = ws.cell(row=2, column=flag_col,
                   value="↑ 期間を選ぶと範囲外の行が薄くなります / "
                         "pick a period above - rows outside it turn grey")
    note.font = NOTE_FONT
    body = f"A{data_start}:{get_column_letter(flag_col)}{last_row}"
    ws.conditional_formatting.add(body, FormulaRule(
        formula=[f"${letter}{data_start}=0"], font=OUT_OF_RANGE_FONT, stopIfTrue=False))
    ws["A2"].comment = None


def _matrix_sheet(
    wb: Workbook,
    title: str,
    matrix: pd.DataFrame,
    subtitle: str,
    diverging: bool,
    zero_rule: Optional[Rule],
    round_to: int = 3,
    style: Optional[dict[str, Any]] = None,
    bar_column: bool = False,
) -> Worksheet:
    ws = wb.create_sheet(title[:31])
    ws["A1"] = subtitle
    ws["A1"].font = TITLE_FONT
    body = matrix.reset_index()
    body.columns = ["Time"] + [str(c) for c in matrix.columns]
    nfmt = {c: ("0.000" if round_to >= 3 else "0.00") for c in body.columns[1:]}
    data_start = _write_table(ws, body, list(body.columns), start_row=3,
                              time_column="Time",
                              number_formats={**nfmt, "Time": TIME_FMT},
                              freeze=True, autofilter=True)
    last_row = data_start + len(body) - 1
    flag_col = len(body.columns) + 1          # one past the data
    if last_row >= data_start:
        _period_picker(ws, wb, matrix, data_start, last_row, flag_col)
    if last_row >= data_start and len(body.columns) > 1:
        rng = (f"B{data_start}:"
               f"{get_column_letter(len(body.columns))}{last_row}")
        st = style or {}
        # The zero rule goes first and stops there: Excel applies rules in
        # priority order, so a band rule added earlier would repaint the zeros.
        if zero_rule is not None and zero_rule.enabled:
            for xr in excel_rules(zero_rule):
                try:
                    xr.stopIfTrue = True
                except AttributeError:
                    pass
                ws.conditional_formatting.add(rng, xr)
        if st.get("mode") != "off":
            for xr in _gradient_cf(st, diverging):
                ws.conditional_formatting.add(rng, xr)
        if st.get("bars") and not diverging:
            from openpyxl.formatting.rule import DataBarRule
            ws.conditional_formatting.add(rng, DataBarRule(
                start_type="min", end_type="max",
                color=excel_argb(st.get("bar_color") or "#4A90D9")[2:],
                showValue=True))
    return ws


def write_analysis_workbook(
    out_path: Path,
    result: AnalysisResult,
    rules: list[Rule],
    opts: OutputOptions,
    info: dict[str, Any],
    png_paths: Optional[dict[str, Path]] = None,
    custom_charts: Optional[list[tuple[str, Path]]] = None,
    analysis: Optional[Any] = None,      # AnalysisOptions: the feature toggles
) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    ws["A1"] = "DataScope - cross-unit analysis"
    ws["A1"].font = Font(bold=True, size=14, color="FF1F3A5F")
    r = 3
    for k, v in info.items():
        ws.cell(row=r, column=1, value=str(k)).font = Font(bold=True, size=10)
        ws.cell(row=r, column=2, value=("" if v is None else str(v))).font = NOTE_FONT
        r += 1
    # the colours used below come from the app's own colour-scale rule
    style = gradient_style(rules)
    ws.cell(row=r, column=1, value="Table colours").font = Font(bold=True, size=10)
    ws.cell(row=r, column=2, value=(
        f"{style.get('label') or 'colour scale'} · {style.get('mode')} · {style.get('bounds')}"
        + (f" · steps {', '.join(str(s) for s in style.get('stops') or [])}"
           if style.get("bounds") == "numbers" and style.get("stops") else "")
        + ("" if style.get("from_rule")
           else "  (no colour-scale rule enabled in the app - default ramp)")
    )).font = NOTE_FONT
    r += 1
    pal = [excel_argb(c) for c in (style.get("colors") or ["FFFCFCFB", "FF184F95"])]
    lo_c, hi_c = pal[0], pal[-1]

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 80

    diag0 = getattr(result, "diagnostics", None)
    if (diag0 is not None and diag0.findings and analysis is not None
            and analysis.wants("findings", "excel")):
        r += 2
        ws.cell(row=r, column=1, value="Findings / 所見").font = TITLE_FONT
        r += 1
        for line in diag0.findings:
            ws.cell(row=r, column=1, value="•")
            ws.cell(row=r, column=2, value=line).font = Font(size=10)
            r += 1
        for note in diag0.notes:
            ws.cell(row=r, column=2, value=note).font = NOTE_FONT
            r += 1

    if result.quality is not None and not result.quality.empty:
        r += 2
        ws.cell(row=r, column=1, value="Source files").font = TITLE_FONT
        _write_table(ws, result.quality, list(result.quality.columns), start_row=r + 1,
                     freeze=False, autofilter=False)

    zero_rule = next((x for x in rules if x.kind is RuleKind.ZERO and x.enabled), None)

    used_names: set[str] = set()

    def uniq(base: str) -> str:
        name = base[:31]
        i = 2
        while name in used_names:
            suffix = f"~{i}"
            name = base[: 31 - len(suffix)] + suffix
            i += 1
        used_names.add(name)
        return name

    for mi, m in enumerate(result.matrices, start=1):
        unit = f" [{m.unit}]" if m.unit else ""
        vs = _matrix_sheet(
            wb, uniq(f"{mi:02d} {m.title}"), m.values,
            f"{m.title}{unit} - value by {'/'.join(m.entities)} at the same timestamp",
            diverging=False, zero_rule=zero_rule, round_to=3, style=style,
        )
        if m.deviation is not None and not m.deviation.empty:
            suffix = "%" if m.deviation_percent else unit
            _matrix_sheet(
                wb, uniq(f"{mi:02d} {m.title} dev"), m.deviation,
                f"{m.title} - deviation from fleet {m.reference_mode} ({suffix})",
                diverging=True, zero_rule=None, round_to=3, style=style,
            )
        # native chart from a hidden downsampled sheet
        if opts.excel_charts:
            small = _downsample(m.values)
            wsd = wb.create_sheet(uniq(f"_cd{mi:02d}"))
            wsd.sheet_state = "hidden"
            wsd.append(["Time"] + [str(c) for c in small.columns])
            for idx, row in small.iterrows():
                wsd.append([idx] + [None if pd.isna(v) else float(v) for v in row.to_numpy()])
            wsc = wb.create_sheet(uniq(f"{mi:02d} {m.title} chart"))
            _line_chart(wsd, wsc, "A2", f"{m.title}{unit} - all units",
                        m.unit, min_col=2, max_col=1 + len(small.columns),
                        min_row=2, max_row=wsd.max_row, width=34, height=14)
            if m.deviation is not None and not m.deviation.empty:
                dsmall = _downsample(m.deviation)
                wsd2 = wb.create_sheet(uniq(f"_dd{mi:02d}"))
                wsd2.sheet_state = "hidden"
                wsd2.append(["Time"] + [str(c) for c in dsmall.columns])
                for idx, row in dsmall.iterrows():
                    wsd2.append([idx] + [None if pd.isna(v) else float(v)
                                         for v in row.to_numpy()])
                _line_chart(wsd2, wsc, "A32",
                            f"{m.title} - deviation from {m.reference_mode}"
                            f" ({'%' if m.deviation_percent else m.unit})",
                            "%" if m.deviation_percent else m.unit,
                            min_col=2, max_col=1 + len(dsmall.columns),
                            min_row=2, max_row=wsd2.max_row, width=34, height=14)
            if png_paths and m.metric in png_paths:
                try:
                    wsc.add_image(XLImage(str(png_paths[m.metric])), "T2")
                except Exception:
                    pass

    if result.daily is not None and not result.daily.empty:
        wsd = wb.create_sheet("Daily summary")
        start = _write_table(wsd, result.daily, list(result.daily.columns), start_row=1,
                             freeze=True, autofilter=True)
        last = start + len(result.daily) - 1
        # Coverage: more is better -> high colour at the top.
        # Zero %: more is worse   -> the scale is flipped on purpose.
        for name, flip in (("Coverage %", False), ("Zero %", True),
                           ("Total", False), ("Mean", False)):
            if name in result.daily.columns:
                j = list(result.daily.columns).index(name) + 1
                letter = get_column_letter(j)
                wsd.conditional_formatting.add(
                    f"{letter}{start}:{letter}{last}",
                    ColorScaleRule(start_type="min", start_color=hi_c if flip else lo_c,
                                   end_type="max", end_color=lo_c if flip else hi_c))

    if custom_charts:
        wsu = wb.create_sheet("User charts")
        wsu["A1"] = "Charts defined on the Analysis tab"
        wsu["A1"].font = TITLE_FONT
        row = 3
        for title, path in custom_charts:
            wsu.cell(row=row, column=1, value=title).font = Font(bold=True, size=10)
            try:
                wsu.add_image(XLImage(str(path)), f"A{row + 1}")
            except Exception:
                wsu.cell(row=row + 1, column=1,
                         value=f"(image not embedded: {path.name})").font = NOTE_FONT
            row += 26
        wsu.column_dimensions["A"].width = 40

    if result.alarms is not None and not result.alarms.empty:
        wsa = wb.create_sheet("Status alarms")
        start = _write_table(wsa, result.alarms, list(result.alarms.columns),
                             start_row=1, freeze=True, autofilter=True)
        last = start + len(result.alarms) - 1
        if "Share %" in result.alarms.columns:
            j = list(result.alarms.columns).index("Share %") + 1
            letter = get_column_letter(j)
            wsa.conditional_formatting.add(
                f"{letter}{start}:{letter}{last}",
                ColorScaleRule(start_type="num", start_value=0, start_color=lo_c,
                               end_type="num", end_value=100, end_color=hi_c))

    # ---- diagnostics, each only if it was targeted at Excel --------------- #
    diag = getattr(result, "diagnostics", None)
    if diag is not None and analysis is not None:
        def _diag_sheet(title: str, df, note: str = "",
                        good_cols: tuple[str, ...] = (),
                        bad_cols: tuple[str, ...] = ()) -> None:
            if df is None or getattr(df, "empty", True):
                return
            ws2 = wb.create_sheet(title[:31])
            row = 1
            if note:
                ws2.cell(row=1, column=1, value=note).font = NOTE_FONT
                row = 3
            start = _write_table(ws2, df, [str(c) for c in df.columns],
                                 start_row=row, freeze=True, autofilter=True)
            last = start + len(df) - 1
            for name, flip in ([(c, False) for c in good_cols]
                               + [(c, True) for c in bad_cols]):
                if name in df.columns:
                    j = list(df.columns).index(name) + 1
                    letter = get_column_letter(j)
                    ws2.conditional_formatting.add(
                        f"{letter}{start}:{letter}{last}",
                        ColorScaleRule(start_type="min", start_color=hi_c if flip else lo_c,
                                       end_type="max", end_color=lo_c if flip else hi_c))

        if analysis.wants("attribution", "excel"):
            _diag_sheet(
                "Attribution", diag.attribution,
                note="Where each unit's shortfall is. "
                     "DC side low = array/strings/shading/soiling; "
                     "efficiency low = the inverter itself. / "
                     "不足の所在：直流側が低い＝アレイ側、変換効率が低い＝インバータ。",
                good_cols=("AC vs fleet %", "DC side vs fleet %",
                           "Efficiency vs fleet pt"))
        if analysis.wants("availability", "excel"):
            _diag_sheet(
                "Availability", diag.availability,
                note="Minutes with usable light but no output, priced against the "
                     "fleet median unit in the same minutes. / "
                     "日射があるのに出力がない時間と、その損失量。",
                good_cols=("Availability %",), bad_cols=("Lost kWh vs fleet",))
        if analysis.wants("efficiency", "excel"):
            _diag_sheet("Efficiency", diag.efficiency,
                        note="DC → AC conversion efficiency, median over the "
                             "compared minutes. / 直流→交流 変換効率（中央値）。",
                        good_cols=("Efficiency %", "Efficiency vs fleet pt"))
        if analysis.wants("performance", "excel"):
            _diag_sheet("Performance ratio", diag.performance,
                        note="PR = produced energy / (irradiation x DC nameplate). / "
                             "PR ＝ 発電量 ÷（日射量 × 直流定格）。",
                        good_cols=("PR",))
        if analysis.wants("tidy", "excel") and diag.tidy is not None \
                and not diag.tidy.empty:
            wst = wb.create_sheet("Data (long)")
            wst.cell(row=1, column=1,
                     value="One row per time/unit/metric - insert a PivotTable on "
                           "this sheet. / 1行＝時刻×号機×項目。ピボットテーブル用。"
                     ).font = NOTE_FONT
            _write_table(wst, diag.tidy, [str(c) for c in diag.tidy.columns],
                         start_row=3, time_column="Time", freeze=True, autofilter=True)

    if result.zero_report is not None and not result.zero_report.empty:
        wsz = wb.create_sheet("Zero report")
        start = _write_table(wsz, result.zero_report, list(result.zero_report.columns),
                             start_row=1, freeze=True, autofilter=True)
        last = start + len(result.zero_report) - 1
        j = list(result.zero_report.columns).index("Zero %") + 1
        letter = get_column_letter(j)
        # many zeros is bad news, so this column is deliberately reversed
        wsz.conditional_formatting.add(
            f"{letter}{start}:{letter}{last}",
            ColorScaleRule(start_type="num", start_value=0, start_color=hi_c,
                           end_type="num", end_value=100, end_color=lo_c))

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
